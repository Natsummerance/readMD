"""Generated document fixtures exercise conversion and real exported artifacts."""
import re
import zipfile
from pathlib import Path
from urllib.parse import unquote
from xml.etree import ElementTree as ET

import fitz
import pytest
from PIL import Image
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches

from src.readmd_modules import convert, mdexport, ocr
from src.readmd_modules.rich_documents import epub_to_md


@pytest.fixture
def picture(tmp_path):
    target = tmp_path / 'test image.png'
    Image.new('RGB', (200, 800), '#4477aa').save(target)
    return target


def assert_assets(markdown, directory):
    sources = re.findall(r'!\[[^\]]*\]\(([^)]+)\)', markdown)
    assert sources, markdown
    assert all((directory / unquote(src)).is_file() for src in sources)


def test_word_content_controls_revisions_tables_and_images(tmp_path, picture):
    document = Document()
    document.add_heading('Word fixture', 1)
    document.add_picture(str(picture), width=Inches(1))
    control = OxmlElement('w:sdt')
    content = OxmlElement('w:sdtContent')
    control.append(content)
    paragraph = document.add_paragraph('Content control text')
    content.append(paragraph._p)
    document.element.body.insert(1, control)
    paragraph = document.add_paragraph()
    inserted = OxmlElement('w:ins')
    run = OxmlElement('w:r')
    text = OxmlElement('w:t'); text.text = 'Accepted revision'
    run.append(text); inserted.append(run); paragraph._p.append(inserted)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Name'; table.cell(0, 1).text = 'Value'
    table.cell(1, 0).text = 'Item'; table.cell(1, 1).text = '42'
    source = tmp_path / 'complex.docx'; document.save(source)
    markdown = convert.docx2md(str(source), form_tables=False)
    assert all(s in markdown for s in ['Word fixture', 'Content control text', 'Accepted revision', '42'])
    assert '|' in markdown
    assert_assets(markdown, tmp_path)


def test_powerpoint_images_tables_charts_notes(tmp_path, picture):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches as PtInches
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = 'Quarterly report'
    slide.shapes.add_picture(str(picture), PtInches(1), PtInches(1), height=PtInches(2))
    table = slide.shapes.add_table(2, 2, PtInches(3), PtInches(1), PtInches(3), PtInches(1)).table
    table.cell(0, 0).text = 'Metric'; table.cell(0, 1).text = 'Result'
    table.cell(1, 0).text = 'Users'; table.cell(1, 1).text = '120'
    data = CategoryChartData(); data.categories = ['Alpha', 'Beta']; data.add_series('Revenue', (10, 20))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, PtInches(3), PtInches(3), PtInches(4), PtInches(2), data)
    slide.notes_slide.notes_text_frame.text = 'Speaker note retained'
    source = tmp_path / 'report.pptx'; presentation.save(source)
    markdown = convert._pptx_to_md(str(source))
    assert all(s in markdown for s in ['Quarterly report', 'Users', '120', 'Revenue', 'Alpha', 'Speaker note retained'])
    assert_assets(markdown, tmp_path)


def test_epub_spine_order_rich_content_and_images(tmp_path, picture):
    source = tmp_path / 'book.epub'
    with zipfile.ZipFile(source, 'w') as archive:
        archive.writestr('META-INF/container.xml', '<container><rootfiles><rootfile full-path="OPS/book.opf"/></rootfiles></container>')
        archive.writestr('OPS/book.opf', '<package><manifest><item id="a" href="a.xhtml"/><item id="z" href="z.xhtml"/></manifest><spine><itemref idref="z"/><itemref idref="a"/></spine></package>')
        archive.writestr('OPS/z.xhtml', '<html><body><h1>First chapter</h1><p><strong>Bold</strong> <a href="https://example.com">Link</a></p><img src="images/pic.png" alt="Figure"/><table><tr><th>Key</th><th>Value</th></tr><tr><td>Answer</td><td>42</td></tr></table></body></html>')
        archive.writestr('OPS/a.xhtml', '<html><body><h1>Second chapter</h1></body></html>')
        archive.write(picture, 'OPS/images/pic.png')
    markdown = epub_to_md(source)
    assert markdown.index('First chapter') < markdown.index('Second chapter')
    assert '**Bold**' in markdown and 'https://example.com' in markdown and '42' in markdown
    assert_assets(markdown, tmp_path)


def test_pdf_columns_and_scanned_page_fallback(tmp_path, picture, monkeypatch):
    source = tmp_path / 'columns.pdf'
    with fitz.open() as document:
        page = document.new_page(width=600, height=800)
        # Alternating creation order must not interleave column reading order.
        for index in range(5):
            page.insert_text((40, 80 + 40 * index), 'LEFT_%d paragraph content retained.' % index, fontsize=10)
            page.insert_text((325, 80 + 40 * index), 'RIGHT_%d paragraph content retained.' % index, fontsize=10)
        page = document.new_page(width=600, height=800)
        page.insert_image(fitz.Rect(100, 100, 300, 700), filename=str(picture))
        document.save(source)
    monkeypatch.setattr(ocr, 'ocr_image', lambda *args, **kwargs: '')
    markdown = convert.pdf2md(str(source))
    assert markdown.index('LEFT_4') < markdown.index('RIGHT_0')
    assert 'Page 2' in markdown and 'OCR unavailable' in markdown
    assert_assets(markdown, tmp_path)


OPTIONS = {
    'page': {'size': 'Custom', 'width': 180, 'height': 240, 'marginLeft': 15, 'marginRight': 15},
    'typography': {'font': 'Arial', 'size': 11, 'firstLineIndent': 5, 'lineHeight': 1.7},
    'header': {'text': 'ReadMD layout verification', 'align': 'center'},
    'images': {'widthPct': 60, 'maxHeightPct': 60},
    'toc': {'enabled': True},
}


def sample_markdown(picture):
    return '# Layout sample\n\nBody with **bold** and [link](https://example.com).\n\n![Tall figure](test%20image.png)\n\n| Name | Value |\n| --- | --- |\n| Item | 42 |\n\n<!-- pagebreak -->\n\n## Second section\n\nEnd marker.\n'


def test_docx_layout_and_native_features(tmp_path, picture):
    target = tmp_path / 'layout.docx'
    result = mdexport.export('docx', sample_markdown(picture), str(tmp_path), target, OPTIONS)
    assert result['ok'], result
    document = Document(target); section = document.sections[0]
    assert section.page_width.mm == pytest.approx(180, abs=.1)
    assert section.page_height.mm == pytest.approx(240, abs=.1)
    assert section.header.paragraphs[0].text == OPTIONS['header']['text']
    assert document.styles['Normal'].paragraph_format.first_line_indent.mm == pytest.approx(5, abs=.1)
    assert document.inline_shapes[0].height.mm <= 121
    assert all(run.font.name == 'Arial' for row in document.tables[0].rows for cell in row.cells for p in cell.paragraphs for run in p.runs)
    with zipfile.ZipFile(target) as archive:
        xml = archive.read('word/document.xml').decode()
        assert 'tblHeader' in xml and 'TOC ' in xml and 'w:type="page"' in xml


def test_pdf_custom_layout_long_code_and_table(tmp_path, picture):
    target = tmp_path / 'layout.pdf'
    content = sample_markdown(picture)
    content += '\n```text\n' + '\n'.join('line %03d <literal> & value' % index for index in range(180)) + '\n```\n'
    content += '\n| Long cell | Result |\n| --- | --- |\n| ' + ('wrapped cell content ' * 320) + 'TABLE_END | done |\n\nDOCUMENT_END\n'
    result = mdexport.export('pdf', content, str(tmp_path), target, OPTIONS)
    assert result['ok'], result
    with fitz.open(target) as document:
        text = '\n'.join(page.get_text() for page in document)
        assert document.page_count > 4
        assert document[0].rect.width == pytest.approx(180 / 25.4 * 72, abs=.1)
        assert document[0].rect.height == pytest.approx(240 / 25.4 * 72, abs=.1)
        assert all(marker in text for marker in ['ReadMD layout verification', 'line 179', 'TABLE_END', 'DOCUMENT_END'])
        for page in document:
            for block in page.get_text('blocks'):
                assert block[2] <= page.rect.width + 1, block


@pytest.mark.parametrize('format', ['html', 'epub'])
def test_portable_export_with_embedded_assets(tmp_path, picture, format):
    target = tmp_path / ('portable.' + format)
    result = mdexport.export(format, sample_markdown(picture), str(tmp_path), target, OPTIONS)
    assert result['ok'], result
    if format == 'html':
        text = target.read_text(encoding='utf-8')
        assert 'data:image/png;base64,' in text and '180mm 240mm' in text
        assert not any('marked.min.js' in warning for warning in result['warns'])
    else:
        with zipfile.ZipFile(target) as archive:
            assert any(name.startswith('OEBPS/images/') for name in archive.namelist())
            chapters = [name for name in archive.namelist() if name.endswith('.xhtml')]
            for chapter in chapters: ET.fromstring(archive.read(chapter))
            assert any(b'<table>' in archive.read(name) and b'<img ' in archive.read(name) for name in chapters)
