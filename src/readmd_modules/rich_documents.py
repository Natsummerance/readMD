"""Structure-preserving document adapters and deterministic extracted assets."""
import hashlib
import html
import os
import posixpath
import re
import zipfile
from pathlib import Path
from urllib.parse import quote, unquote, urlsplit
from xml.etree import ElementTree as ET


def save_asset(source, data, extension):
    if len(data) > 32 * 1024 * 1024:
        raise ValueError('document_image_too_large')
    extension = str(extension).lower().lstrip('.')
    if extension not in {'png', 'jpg', 'jpeg', 'gif', 'webp', 'svg', 'bmp', 'tif', 'tiff', 'emf', 'wmf'}:
        extension = 'bin'
    root = Path(source).resolve()
    directory = root.with_name(root.stem + '.assets')
    directory.mkdir(exist_ok=True)
    name = hashlib.sha256(data).hexdigest()[:24] + '.' + extension
    target = directory / name
    if not target.exists():
        try:
            with target.open('xb') as handle:
                handle.write(data)
        except FileExistsError:
            pass
    return quote(directory.name + '/' + name, safe='/')


def _member(archive, name):
    info = archive.getinfo(name)
    if info.file_size > 32 * 1024 * 1024:
        raise ValueError('document_member_too_large')
    return archive.read(info)


def epub_to_md(source):
    from bs4 import BeautifulSoup
    from markdownify import markdownify
    with zipfile.ZipFile(source) as archive:
        names = set(archive.namelist())
        if len(names) > 10000 or sum(i.file_size for i in archive.infolist()) > 256 * 1024 * 1024:
            raise ValueError('epub_too_large')
        try:
            container = ET.fromstring(_member(archive, 'META-INF/container.xml'))
            opf = container.find('.//{*}rootfile').get('full-path')
            package = ET.fromstring(_member(archive, opf))
            manifest = {i.get('id'): posixpath.normpath(posixpath.join(posixpath.dirname(opf), unquote(i.get('href', '')))) for i in package.findall('.//{*}manifest/{*}item')}
            chapters = [manifest[i.get('idref')] for i in package.findall('.//{*}spine/{*}itemref') if i.get('idref') in manifest and i.get('linear') != 'no']
        except (KeyError, AttributeError, ET.ParseError):
            chapters = sorted(n for n in names if n.lower().endswith(('.xhtml', '.html', '.htm')))
        chunks = []
        for chapter in chapters:
            if chapter not in names:
                continue
            soup = BeautifulSoup(_member(archive, chapter), 'html.parser')
            for unwanted in soup(['script', 'style', 'noscript']):
                unwanted.decompose()
            for image in soup.find_all('img'):
                src = urlsplit(image.get('src', ''))
                if src.scheme or src.netloc:
                    continue
                member = posixpath.normpath(posixpath.join(posixpath.dirname(chapter), unquote(src.path)))
                if member in names:
                    image['src'] = save_asset(source, _member(archive, member), Path(member).suffix)
            for link in soup.find_all('a', href=True):
                href = urlsplit(link['href'])
                if not href.scheme and href.fragment:
                    link['href'] = '#' + href.fragment
            text = markdownify(str(soup.body or soup), heading_style='ATX', bullets='-').strip()
            if text:
                chunks.append(text)
        if not chunks:
            raise ValueError('epub-empty')
        return '\n\n---\n\n'.join(chunks)


def pptx_to_md(source):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    presentation = Presentation(source)
    output = ['# ' + Path(source).stem]

    def cell(value):
        return str(value).replace('|', '\\|').replace('\n', '<br>')

    def shape_lines(shapes, title):
        lines = []
        for shape in sorted(shapes, key=lambda item: (item.top, item.left)):
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                lines.extend(shape_lines(shape.shapes, title))
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image
                lines.append('![%s](%s)' % (cell(shape.name), save_asset(source, image.blob, image.ext)))
            elif shape.has_table:
                rows = [[cell(c.text) if not c.is_spanned else '' for c in row.cells] for row in shape.table.rows]
                if rows:
                    lines += ['| ' + ' | '.join(rows[0]) + ' |', '| ' + ' | '.join(['---'] * len(rows[0])) + ' |']
                    lines += ['| ' + ' | '.join(row) + ' |' for row in rows[1:]]
            elif shape.has_chart:
                chart = shape.chart
                try:
                    categories = [c.label for c in chart.plots[0].categories]
                    series = list(chart.series)
                    lines += ['| Category | ' + ' | '.join(cell(s.name) for s in series) + ' |', '| --- | ' + ' | '.join('---' for _ in series) + ' |']
                    for i, category in enumerate(categories):
                        lines.append('| ' + cell(category) + ' | ' + ' | '.join(cell(s.values[i]) if i < len(s.values) else '' for s in series) + ' |')
                except (AttributeError, IndexError, ValueError):
                    lines.append('> ' + shape.name)
            elif shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = ''
                    for run in paragraph.runs:
                        value = run.text
                        if run.font.bold and value.strip(): value = '**' + value + '**'
                        if run.font.italic and value.strip(): value = '*' + value + '*'
                        if run.hyperlink.address: value = '[' + value + '](' + run.hyperlink.address + ')'
                        text += value
                    if not text.strip(): continue
                    if shape == title:
                        lines.append('## ' + text)
                    elif paragraph.level or paragraph._p.find('.//{*}buChar') is not None:
                        lines.append('  ' * paragraph.level + '- ' + text)
                    else:
                        lines.append(text)
            lines.append('')
        return lines

    for number, slide in enumerate(presentation.slides, 1):
        if not slide.shapes.title:
            output.append('## Slide %d' % number)
        output.extend(shape_lines(slide.shapes, slide.shapes.title))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                output.extend(['### Notes', notes, ''])
    return '\n'.join(output).strip() + '\n'


def word_children(element):
    """Read accepted revisions and content controls without flattening tables."""
    for child in element:
        local = child.tag.rsplit('}', 1)[-1] if isinstance(child.tag, str) else ''
        if local in {'sdt', 'sdtContent', 'ins', 'smartTag', 'customXml'}:
            yield from word_children(child)
        elif local not in {'del', 'sdtPr'}:
            yield child


def pdf_columns(page, render, body_size):
    """Split only pages with a clear central gutter; keep full-width headings."""
    import fitz
    # MuPDF can group both columns into one block; line boxes retain the gutter.
    blocks = [(*line['bbox'], ''.join(span.get('text', '') for span in line['spans']))
              for block in page.get_text('dict').get('blocks', []) if block.get('type') == 0
              for line in block.get('lines', []) if line.get('spans')]
    middle = (page.rect.x0 + page.rect.x1) / 2
    left = [b for b in blocks if b[2] < middle - 6]
    right = [b for b in blocks if b[0] > middle + 6]
    wide = [b for b in blocks if b not in left and b not in right]
    if sum(len(b[4]) for b in left) < 80 or sum(len(b[4]) for b in right) < 80:
        return None
    if any(b[2] - b[0] < page.rect.width * .55 for b in wide):
        return None
    try:
        if any(table.bbox[0] < middle < table.bbox[2] for table in page.find_tables().tables):
            return None
    except (AttributeError, ValueError):
        pass

    class Clip:
        def __init__(self, rect): self.rect = rect
        def get_text(self, *args, **kwargs): return page.get_text(*args, **{**kwargs, 'clip': self.rect})
        def find_tables(self, *args, **kwargs): return page.find_tables(*args, **{**kwargs, 'clip': self.rect})
        def __getattr__(self, name): return getattr(page, name)

    output, top = [], page.rect.y0
    for block in sorted(wide, key=lambda b: b[1]) + [None]:
        bottom = block[1] if block else page.rect.y1
        if bottom > top:
            for x0, x1 in [(page.rect.x0, middle), (middle, page.rect.x1)]:
                output.append(render(Clip(fitz.Rect(x0, top, x1, bottom)), default_body_size=body_size))
        if block:
            output.append(render(Clip(fitz.Rect(page.rect.x0, block[1], page.rect.x1, block[3])), default_body_size=body_size))
            top = max(top, block[3])
    return '\n\n'.join(text for text in output if text.strip())
